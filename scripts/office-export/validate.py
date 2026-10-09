#!/usr/bin/env python3
"""Independent package and source-IR checks; deliberately never approve delivery."""
from __future__ import annotations

import argparse
from collections import defaultdict
import hashlib
import json
from pathlib import Path
import posixpath
import re
from zipfile import ZipFile

from lxml import etree
from pptx import Presentation
from docx import Document

from generate import ROOT, text_of, safe_asset

NS = {
    'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main',
    'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships',
}


def package_errors(file: Path) -> list[str]:
    errors = []
    with ZipFile(file) as archive:
        names = set(archive.namelist())
        if archive.testzip(): errors.append(f'{file.name}: CRC error')
        for name in names:
            if name.startswith('/') or '..' in name.split('/'):
                errors.append(f'{file.name}: unsafe package path {name}')
            if not name.endswith(('.xml', '.rels')): continue
            try:
                tree = etree.fromstring(archive.read(name), etree.XMLParser(resolve_entities=False, no_network=True))
            except etree.XMLSyntaxError as exc:
                errors.append(f'{file.name}: invalid XML {name}: {exc}'); continue
            if not name.endswith('.rels'): continue
            parent = '' if name == '_rels/.rels' else posixpath.dirname(posixpath.dirname(name))
            for rel in tree:
                if rel.get('TargetMode') == 'External': continue
                target = rel.get('Target', '')
                normalized = posixpath.normpath(posixpath.join(parent, target)).lstrip('/')
                if normalized not in names:
                    errors.append(f'{file.name}: dangling relationship {name} -> {target}')
    return errors


def token_text(runs):
    return ''.join(f'⟦{run["asset"]}⟧' if 'asset' in run else run.get('text', '') for run in runs)


def expected_block(block):
    if block['type'] in ('heading', 'paragraph'):
        return block.get('prefix', '') + token_text(block['runs'])
    if block['type'] == 'code': return block['text']
    if block['type'] == 'graphic': return f'⟦{block["asset"]}⟧'
    return ''


def word_text(element):
    values = []
    for child in element.iter():
        if child.tag == f'{{{NS["w"]}}}t': values.append(child.text or '')
        elif child.tag == f'{{{NS["w"]}}}tab': values.append('\t')
        elif child.tag in (f'{{{NS["w"]}}}br', f'{{{NS["w"]}}}cr'): values.append('\n')
    return ''.join(values)


def validate(directory: Path) -> dict:
    source = json.loads((directory / 'source.json').read_text())
    plan = json.loads((directory / 'layout-plan.json').read_text())
    errors, checks = [], []
    if hashlib.sha256((directory / 'source.md').read_bytes()).hexdigest() != source['source_sha256']:
        errors.append('Original Markdown snapshot hash changed')
    for kind, digest in plan['template_sha256'].items():
        if hashlib.sha256((ROOT / plan['profile'][kind]['path']).read_bytes()).hexdigest() != digest:
            errors.append(f'Original {kind} template changed since generation')
    for ext in ('pptx', 'docx'):
        errors.extend(package_errors(directory / f'aibs-markdown-proof.{ext}'))
    checks.append('ZIP CRC / XML parsing / internal relationship targets / original template hashes')

    prs = Presentation(directory / 'aibs-markdown-proof.pptx')
    if len(prs.slides) != len(plan['pages']): errors.append('Slide count does not match plan')
    reconstructed = defaultdict(str)
    table_rows = defaultdict(list)
    seen_headers = set()
    assets = {asset['id']: asset for asset in source['assets']}
    actual_ppt_links = set()
    for slide, page in zip(prs.slides, plan['pages']):
        if page.get('notes'):
            expected_notes = '\n\n'.join(f'Markdown source [{item["block"]}]\n{item["text"]}' for item in page['notes'])
            if not slide.has_notes_slide or slide.notes_slide.notes_text_frame.text != expected_notes:
                errors.append(f'Mermaid speaker-note source differs on slide: {page["title"]}')
            else:
                for item in page['notes']: reconstructed[item['block']] += item['text']
        shapes = {shape.name: shape for shape in slide.shapes}
        for rel in slide.part.rels.values():
            if rel.is_external: actual_ppt_links.add(rel.target_ref)
        for element in page['elements']:
            shape = shapes.get(element['name'])
            if shape is None: errors.append(f'Missing PPT shape: {element["name"]}'); continue
            block = element.get('block')
            if shape.left < 0 or shape.top < 0 or shape.left + shape.width > prs.slide_width + 12700 or shape.top + shape.height > prs.slide_height + 12700:
                errors.append(f'Shape outside slide: {shape.name}')
            for key, actual in [('x', shape.left), ('y', shape.top), ('w', shape.width), ('h', shape.height)]:
                if abs(actual / 12700 - element[key]) > .15: errors.append(f'Geometry differs: {shape.name}/{key}')
            if element['kind'] == 'text':
                actual = shape.text.replace('\v', '\n')
                original = text_of(element['runs'])
                expected = '\n'.join(text_of(line) for line in element['line_runs']) if 'line_runs' in element else original
                if actual != expected: errors.append(f'PPT text differs: {shape.name}')
                if 'line_runs' in element and ''.join(text_of(line) for line in element['line_runs']) != original.replace('\n', ''):
                    errors.append(f'Heading text lost during line layout: {shape.name}')
                if block: reconstructed[block] += original if 'line_runs' in element and actual == expected else actual
            elif element['kind'] == 'image':
                asset = assets[element['asset']]
                if shape.image.blob != safe_asset(directory, asset['file']).read_bytes(): errors.append(f'PPT image differs: {shape.name}')
                if shape._element.nvPicPr.cNvPr.get('descr') != asset['alt']: errors.append(f'PPT alt text missing: {shape.name}')
                if block: reconstructed[block] += f'⟦{asset["id"]}⟧'
            elif element['kind'] == 'table':
                for i, row in enumerate(element['rows']):
                    actual = [cell.text.replace('\v', '\n') for cell in shape.table.rows[i].cells]
                    if actual != [text_of(cell) for cell in row]: errors.append(f'PPT table content differs: {shape.name}, row {i}')
                    source_index = element['row_indices'][i]
                    if source_index == 0 and block in seen_headers: continue
                    if source_index == 0: seen_headers.add(block)
                    table_rows[block].append((source_index, actual))
    for block in source['blocks']:
        if block['type'] == 'table':
            expected = [(i, [text_of(cell) for cell in row]) for i, row in enumerate(block['rows'])]
            if table_rows[block['id']] != expected: errors.append(f'Table rows lost/reordered: {block["id"]}')
        elif block['type'] not in ('rule', 'break') and reconstructed[block['id']] != expected_block(block):
            errors.append(f'PPT source-IR block lost/changed: {block["id"]}')
    checks.append('Editable PPT text, code whitespace, Mermaid speaker-note source, inline-asset order, table row coverage, embedded image bytes and shape bounds')

    doc = Document(directory / 'aibs-markdown-proof.docx')
    actual = word_text(doc.element.body)
    asset_hashes = {asset['id']: hashlib.sha256(safe_asset(directory, asset['file']).read_bytes()).hexdigest() for asset in source['assets']}
    word_tokens = []
    for node in doc.element.body.iter():
        if node.tag == f'{{{NS["w"]}}}t': word_tokens.append(node.text or '')
        elif node.tag == f'{{{NS["w"]}}}tab': word_tokens.append('\t')
        elif node.tag in (f'{{{NS["w"]}}}br', f'{{{NS["w"]}}}cr'): word_tokens.append('\n')
        elif node.tag == '{http://schemas.openxmlformats.org/drawingml/2006/main}blip':
            rid = node.get(f'{{{NS["r"]}}}embed')
            if rid in doc.part.related_parts:
                digest = hashlib.sha256(doc.part.related_parts[rid].blob).hexdigest()
                if digest in asset_hashes.values(): word_tokens.append(f'⟦{digest}⟧')
    normalize = lambda value: re.sub(r'\s+', '', value)
    normalized, cursor = normalize(''.join(word_tokens)), 0
    expected_links = set()
    for block in source['blocks']:
        if block['type'] == 'table':
            groups = [cell for row in block['rows'] for cell in row]
        elif block['type'] in ('paragraph', 'heading'):
            groups = [[{'text': block.get('prefix', '')}, *block['runs']]]
        elif block['type'] == 'code': groups = [[{'text': block['text']}]]
        elif block['type'] == 'graphic': groups = [[{'asset': block['asset']}]]
        else: groups = []
        for group in groups:
            for run in group:
                if run.get('href'): expected_links.add(run['href'])
            value = normalize(''.join(f'⟦{asset_hashes[run["asset"]]}⟧' if 'asset' in run else run.get('text', '') for run in group))
            if not value: continue
            index = normalized.find(value, cursor)
            if index == -1: errors.append(f'Word source-IR text lost/reordered: {block["id"]}')
            else: cursor = index + len(value)
    doc_links = {rel.target_ref for rel in doc.part.rels.values() if rel.is_external}
    if len(doc.tables) != sum(block['type'] == 'table' for block in source['blocks']): errors.append('Word native table count differs')
    for table in doc.tables:
        if not table.rows[0]._tr.xpath('./w:trPr/w:tblHeader'): errors.append('Word repeating table header missing')
    if not expected_links.issubset(doc_links): errors.append('Word hyperlink targets lost')
    if not expected_links.issubset(actual_ppt_links): errors.append('PPT hyperlink targets lost')
    # Exact code whitespace check (normal paragraph whitespace is separately normalized).
    actual_paragraphs = [word_text(p._p) for p in doc.paragraphs]
    for block in source['blocks']:
        if block['type'] == 'code' and block['text'] not in actual_paragraphs:
            errors.append(f'Word code whitespace changed: {block["id"]}')
    alt_texts = [shape._inline.docPr.get('descr') for shape in doc.inline_shapes]
    for asset in source['assets']:
        if asset['alt'] not in alt_texts: errors.append(f'Word graphic/alt missing: {asset["id"]}')
    with ZipFile(directory / 'aibs-markdown-proof.docx') as archive:
        images = [archive.read(name) for name in archive.namelist() if name.startswith('word/media/')]
        for asset in source['assets']:
            if safe_asset(directory, asset['file']).read_bytes() not in images: errors.append(f'Word embedded image differs: {asset["id"]}')
        visible_xml = [name for name in archive.namelist() if name == 'word/document.xml' or re.match(r'word/(header|footer)\d+\.xml$', name)]
        source_text = (directory / 'source.md').read_text()
        for name in visible_xml:
            tree = etree.fromstring(archive.read(name))
            text = word_text(tree)
            for marker in ('SELECT FROM DROPDOWN OPTIONS', '[Statement of Direction Title]', 'PLEASE READ:', 'Lorem ipsum', '使用“开始”选项卡', 'Purpose statement', 'STYLEREF'):
                if marker in text and marker not in source_text: errors.append(f'Stale example/field text in {name}: {marker}')
            fields = tree.xpath('//w:instrText/text() | //w:fldSimple/@w:instr', namespaces=NS)
            if any('STYLEREF' in field or 'TOC ' in field for field in fields): errors.append(f'Stale Word field in {name}')
    checks.append('Word ordered text coverage, exact code whitespace, native tables, hyperlinks, embedded image bytes/alt text, stale fields and example cleanup')
    report = {
        'status': 'failed' if errors else 'structural_checks_passed_pending_office',
        'delivery_approved': False, 'errors': errors, 'checks': checks,
        'counts': {'source_blocks': len(source['blocks']), 'assets': len(assets), 'slides': len(prs.slides), 'word_tables': len(doc.tables)},
        'native_office_acceptance': plan['native_office_acceptance'],
        'limitations': plan['warnings'] + plan['profile']['pending_acceptance'] + [
            'Coverage checks compare the current shared renderer IR, not every possible Markdown dialect.',
            'Text bounding boxes and XML validity do not prove absence of Office text overflow.',
            'PNG graphics are not editable Office charts/equations; text and ordinary tables are native objects.',
        ],
    }
    return report


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('directory', type=Path)
    args = parser.parse_args()
    report = validate(args.directory.resolve())
    (args.directory / 'validation-report.json').write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
    text = ['# AIBS 模板导出样稿检查报告', '', f'状态：`{report["status"]}`', '', '**正式交付验收：未通过（尚未执行 macOS / Windows Office 人工验收）。**', '', '## 自动检查', '']
    text.extend(f'- {check}' for check in report['checks'])
    text.extend(['', '## 数量', '', json.dumps(report['counts'], ensure_ascii=False), '', '## 错误', ''])
    text.extend(f'- {error}' for error in report['errors'])
    if not report['errors']: text.append('- 本轮自动检查未发现错误；不代表 Office 渲染无误。')
    text.extend(['', '## 限制与待验收项', ''])
    text.extend(f'- {value}' for value in report['limitations'])
    (args.directory / 'VALIDATION.md').write_text('\n'.join(text) + '\n', encoding='utf-8')
    print(json.dumps(report, ensure_ascii=False, indent=2))
    raise SystemExit(1 if report['errors'] else 0)


if __name__ == '__main__': main()
