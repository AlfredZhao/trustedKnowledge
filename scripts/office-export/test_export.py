"""Run: python -m unittest discover -s scripts/office-export -p 'test_*.py' -v"""
import hashlib
import json
from pathlib import Path
import tempfile
import unittest

from PIL import Image
from pptx import Presentation
from pptx.enum.text import MSO_AUTO_SIZE
from docx import Document

from generate import HERE, ROOT, Metrics, Planner, safe_asset, slice_runs, text_of, write_pptx, write_docx, load_template_profile
from validate import validate, package_errors


class DeterministicMetrics(Metrics):
    def __init__(self): pass
    def width(self, text, size):
        return sum(1 if ord(char) > 127 else .65 for char in text.expandtabs(4)) * size


class LayoutTests(unittest.TestCase):
    def setUp(self):
        self.profile = load_template_profile()
        self.metrics = DeterministicMetrics()

    def test_slice_runs_preserves_styles_links_and_unicode(self):
        runs = [{'text': '中文😀', 'bold': True}, {'text': 'foo_bar', 'href': 'https://example.com'}]
        parts = [slice_runs(runs, i, i + 2) for i in range(0, len(text_of(runs)), 2)]
        self.assertEqual(''.join(text_of(part) for part in parts), text_of(runs))
        self.assertTrue(parts[0][0]['bold'])
        self.assertEqual(parts[-1][-1]['href'], 'https://example.com')

    def test_slice_never_silently_discards_formula(self):
        with self.assertRaisesRegex(ValueError, 'inline graphic'):
            slice_runs([{'asset': 'a'}], 0, 1)

    def test_code_line_ranges_reconstruct_exact_whitespace(self):
        value = '    keep_indent_and_underscores\n\n' + '中😀\t' * 30 + '\n'
        lines = self.metrics.lines(value, 100, 16)
        self.assertGreater(len(lines), 3)
        self.assertEqual(''.join(value[a:b] for a, b in lines), value)

    def test_unsupported_block_is_an_error(self):
        with self.assertRaisesRegex(ValueError, 'Unsupported block'):
            Planner({'blocks': [{'id': 'x', 'type': 'unknown'}], 'assets': []}, self.profile, self.metrics).build()

    def test_long_title_requires_review_instead_of_clipping(self):
        with self.assertRaisesRegex(ValueError, '标题过长'):
            Planner({'blocks': [{'id': 'x', 'type': 'heading', 'level': 1, 'runs': [{'text': '过长标题' * 100}]}], 'assets': []}, self.profile, self.metrics).build()

    def test_title_layout_preserves_styled_runs_and_identifiers(self):
        runs = [{'text': '揭开 Oracle 属性图的神秘面纱：给 DBA 的 ', 'bold': True}, {'text': 'GRAPH_TABLE', 'href': 'https://example.com'}, {'text': ' 内部机制指南'}]
        source = {'blocks': [{'id': 'title', 'type': 'heading', 'level': 1, 'runs': runs}], 'assets': []}
        elements = Planner(source, self.profile, self.metrics).build()[0]['elements']
        heading = elements[0]
        self.assertEqual(''.join(text_of(line) for line in heading['line_runs']), text_of(runs))
        self.assertEqual(heading['runs'], runs)
        self.assertTrue(any(any(run.get('href') == 'https://example.com' and run['text'] == 'GRAPH_TABLE' for run in line) for line in heading['line_runs']))
        for first, second in zip(elements, elements[1:]):
            self.assertLess(first['y'] + first['h'], second['y'])

    def test_section_titles_and_continuations_stay_above_body(self):
        title = '章节标题 GRAPH_TABLE 机制 ' * 3
        planner = Planner({'assets': [], 'blocks': []}, self.profile, self.metrics)
        planner.new_page(title, 'heading')
        planner.new_page()
        for page in planner.pages:
            item = page['elements'][0]
            self.assertLess(item['y'] + item['h'], planner.box[1])
            self.assertEqual(''.join(text_of(line) for line in item['line_runs']), text_of(item['runs']))
        self.assertTrue(planner.pages[-1]['title'].endswith('（续）'))

    def test_preferred_colon_break_never_rejects_otherwise_fitting_title(self):
        title = '主' * 10 + '：' + '文' * 25
        planner = Planner({'assets': [], 'blocks': []}, self.profile, self.metrics)
        item = planner.heading([{'text': title}], 400, 80, (20,), 2, '章节标题')
        self.assertEqual(len(item['line_runs']), 2)
        self.assertEqual(''.join(text_of(line) for line in item['line_runs']), title)

    def test_oversized_table_row_rejected(self):
        source = {'assets': [], 'blocks': [{'id': 'table', 'type': 'table', 'rows': [[[{'text': '表头'}]], [[{'text': '不能缩成小字' * 1000}]]]}]}
        with self.assertRaisesRegex(ValueError, 'table row'):
            Planner(source, self.profile, self.metrics).build()

    def test_assets_cannot_escape_local_output(self):
        with tempfile.TemporaryDirectory() as tmp:
            directory = Path(tmp)
            (directory / 'secret').write_text('not an asset')
            with self.assertRaisesRegex(ValueError, 'outside'):
                safe_asset(directory, 'secret')


class RoundTripTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temporary = tempfile.TemporaryDirectory()
        cls.directory = Path(cls.temporary.name)
        (cls.directory / 'assets').mkdir()
        Image.new('RGB', (160, 60), 'white').save(cls.directory / 'assets/test.png')
        profile = load_template_profile()
        cls.profile = profile
        cls.template_hashes = {kind: hashlib.sha256((ROOT / profile[kind]['path']).read_bytes()).hexdigest() for kind in ('pptx', 'docx')}
        cls.source = {
            'schema': 1, 'source_sha256': hashlib.sha256(b'proof fixture').hexdigest(),
            'assets': [{'id': 'a1', 'file': 'assets/test.png', 'width': 160, 'height': 60, 'kind': 'formula', 'alt': 'E=mc^2'}],
            'blocks': [
                {'id': 'title', 'type': 'heading', 'level': 1, 'runs': [{'text': '揭开 Oracle 属性图的神秘面纱：给 DBA 的 GRAPH_TABLE 内部机制指南'}]},
                {'id': 'heading', 'type': 'heading', 'level': 2, 'runs': [{'text': '完整内容'}]},
                {'id': 'body', 'type': 'paragraph', 'runs': [{'text': '保真长文。' * 250, 'bold': True}]},
                {'id': 'code', 'type': 'code', 'text': ''.join(f'    line_{i} = {i}\n' for i in range(75))},
                {'id': 'table', 'type': 'table', 'rows': [[[{'text': 'ID'}], [{'text': '内容'}]]] + [[[{'text': str(i)}], [{'text': f'row_{i}'}]] for i in range(36)]},
                {'id': 'inline', 'type': 'paragraph', 'runs': [{'text': '公式前'}, {'asset': 'a1', 'alt': 'E=mc^2'}, {'text': '公式后'}]},
                {'id': 'graphic', 'type': 'graphic', 'asset': 'a1'},
                {'id': 'mermaid-source', 'type': 'code', 'language': 'mermaid', 'text': 'flowchart LR\n  A[Start] --> B[End]'},
                {'id': 'link', 'type': 'paragraph', 'prefix': '7. ', 'runs': [{'text': 'example', 'href': 'https://example.com/docs'}]},
                {'id': 'manual', 'type': 'break'},
                {'id': 'end', 'type': 'paragraph', 'runs': [{'text': 'END-OF-FIXTURE'}]},
            ],
        }
        cls.pages = Planner(cls.source, profile, DeterministicMetrics()).build()
        write_pptx(cls.source, profile, cls.pages, cls.directory)
        write_docx(cls.source, profile, cls.directory)
        (cls.directory / 'source.md').write_bytes(b'proof fixture')
        (cls.directory / 'source.json').write_text(json.dumps(cls.source))
        (cls.directory / 'layout-plan.json').write_text(json.dumps({
            'pages': cls.pages, 'profile': profile, 'template_sha256': cls.template_hashes, 'warnings': [],
            'native_office_acceptance': {'macOS': 'not_run', 'Windows': 'not_run'},
        }))

    @classmethod
    def tearDownClass(cls): cls.temporary.cleanup()

    def test_roundtrip_has_no_structural_errors_but_never_approves_delivery(self):
        report = validate(self.directory)
        self.assertEqual(report['errors'], [])
        self.assertGreater(report['counts']['slides'], 10)
        self.assertFalse(report['delivery_approved'])
        self.assertEqual(report['native_office_acceptance']['macOS'], 'not_run')

    def test_original_templates_remain_byte_identical(self):
        for kind, before in self.template_hashes.items():
            self.assertEqual(hashlib.sha256((ROOT / self.profile[kind]['path']).read_bytes()).hexdigest(), before)

    def test_ppt_table_continuations_have_identical_headers(self):
        parts = [e for page in self.pages for e in page['elements'] if e['kind'] == 'table']
        self.assertGreater(len(parts), 1)
        self.assertTrue(all(e['row_indices'][0] == 0 for e in parts))
        self.assertEqual([i for e in parts for i in e['row_indices'][1:]], list(range(1, 37)))

    def test_no_dangling_package_relationships(self):
        for kind in ('pptx', 'docx'):
            self.assertEqual(package_errors(self.directory / f'aibs-markdown-proof.{kind}'), [])

    def test_fixed_ppt_boxes_do_not_auto_resize_in_office(self):
        prs = Presentation(self.directory / 'aibs-markdown-proof.pptx')
        for slide in prs.slides:
            for shape in slide.shapes:
                if shape.has_text_frame:
                    self.assertEqual(shape.text_frame.auto_size, MSO_AUTO_SIZE.NONE)

    def test_mermaid_source_remains_editable_in_speaker_notes(self):
        prs = Presentation(self.directory / 'aibs-markdown-proof.pptx')
        notes = '\n'.join(slide.notes_slide.notes_text_frame.text for slide in prs.slides if slide.has_notes_slide)
        self.assertIn('flowchart LR\n  A[Start] --> B[End]', notes)

    def test_word_keeps_template_heading_font_and_removes_example_fields(self):
        doc = Document(self.directory / 'aibs-markdown-proof.docx')
        template = Document(ROOT / self.profile['docx']['path'])
        self.assertEqual(doc.styles['Heading 1'].font.name, template.styles['Heading 1'].font.name)
        headings = [p for p in doc.paragraphs if p.style.name.startswith('Heading')]
        self.assertTrue(headings)
        self.assertTrue(all(run.font.name is None for p in headings for run in p.runs))
        self.assertNotIn('STYLEREF', doc.element.xml)
        self.assertEqual(len(doc.tables), 1)

    def test_validator_detects_missing_ppt_text(self):
        file = self.directory / 'aibs-markdown-proof.pptx'
        original = file.read_bytes()
        try:
            prs = Presentation(file)
            target = next(s for slide in prs.slides for s in slide.shapes if s.name.startswith('tk:body:'))
            target.text = 'tampered'
            prs.save(file)
            self.assertTrue(any('text differs' in e for e in validate(self.directory)['errors']))
        finally: file.write_bytes(original)

    def test_validator_detects_changed_code_indentation(self):
        file = self.directory / 'aibs-markdown-proof.docx'
        original = file.read_bytes()
        try:
            doc = Document(file)
            paragraph = next(p for p in doc.paragraphs if p.style.name == 'Code')
            paragraph.text = paragraph.text.replace('    ', '')
            doc.save(file)
            self.assertTrue(any('code whitespace' in e for e in validate(self.directory)['errors']))
        finally: file.write_bytes(original)


if __name__ == '__main__': unittest.main()
