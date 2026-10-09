#!/usr/bin/env python3
"""Generate auditable, editable Office proofs from prepare.mjs's local content IR.

Not a public API or an arbitrary-template importer. Private OOXML operations are
deliberately isolated below and covered by package/content regression checks.
"""
from __future__ import annotations

import argparse
from copy import deepcopy
from datetime import datetime
import hashlib
import html
import json
import re
from pathlib import Path
from zoneinfo import ZoneInfo

from PIL import ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.oxml.xmlchemy import OxmlElement as PptElement
from pptx.util import Inches, Pt
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches as WordInches, Pt as WordPt, RGBColor as WordColor
from docx.opc.constants import RELATIONSHIP_TYPE as REL

ROOT = Path(__file__).resolve().parents[4]
HERE = Path(__file__).resolve().parent


def load_template_profile(now: datetime | None = None) -> dict:
    """Resolve copyright once per job, never at process startup or in body text."""
    zone = ZoneInfo('Asia/Shanghai')
    instant = now if now is not None else datetime.now(zone)
    if instant.tzinfo is None:
        raise ValueError('Copyright year requires a timezone-aware datetime')
    year = str(instant.astimezone(zone).year)
    profile = json.loads((HERE / 'template-profile.json').read_text(encoding='utf-8'))
    for kind in ('pptx', 'docx'):
        profile[kind]['copyright'] = profile[kind]['copyright'].replace('{year}', year)
    return profile


def text_of(runs: list[dict]) -> str:
    return "".join(run.get("text", "") for run in runs)


def safe_asset(directory: Path, relative: str) -> Path:
    file = (directory / relative).resolve()
    if not file.is_relative_to((directory / "assets").resolve()) or not file.is_file():
        raise ValueError(f"Asset is outside the generated assets directory: {relative}")
    return file


def slice_runs(runs: list[dict], start: int, end: int) -> list[dict]:
    """Lossless text slicing; asset-bearing runs must use the inline-graphic path."""
    result, offset = [], 0
    for run in runs:
        if "asset" in run:
            raise ValueError("Cannot silently flatten an inline graphic")
        value = run.get("text", "")
        lo, hi = max(0, start - offset), min(len(value), end - offset)
        if hi > lo:
            result.append({**run, "text": value[lo:hi]})
        offset += len(value)
    return result


class Metrics:
    """Conservative preflight, NOT an Office layout engine or a delivery verdict."""
    def __init__(self, font_file: Path):
        self.file = font_file
        self.font = ImageFont.truetype(str(font_file), 100)

    def width(self, text: str, size: float) -> float:
        # A safety allowance for bold/fallback faces and target-app differences.
        return self.font.getlength(text.expandtabs(4)) * size / 100 * 1.12

    def lines(self, text: str, width: float, size: float) -> list[tuple[int, int]]:
        result, start, used = [], 0, 0.0
        for i, char in enumerate(text):
            if char == "\n":
                result.append((start, i + 1)); start, used = i + 1, 0.0
                continue
            advance = self.width(char, size)
            if advance > width:
                raise ValueError("A character is wider than the available text area")
            if used + advance > width and i > start:
                result.append((start, i)); start, used = i, 0.0
            used += advance
        if start < len(text) or not result:
            result.append((start, len(text)))
        return result


class Planner:
    def __init__(self, source: dict, profile: dict, metrics: Metrics):
        self.source, self.profile, self.metrics = source, profile, metrics
        self.assets = {asset["id"]: asset for asset in source["assets"]}
        self.pages: list[dict] = []
        self.title = "Markdown 文档"
        self.page: dict | None = None
        self.y = 0.0
        self.box = [value * 72 for value in profile["pptx"]["body_inches"]]
        self.sequence = 0
        self.section_pages = 0

    def heading(self, runs, width, height, sizes, max_lines, label, prefer_main_title=True):
        """Explicit line layout shared by PPT and previews; source runs stay intact."""
        text = text_of(runs)
        # A substantial Chinese main-title prefix is a better line boundary than
        # leaving its colon (or final character) at the start of the next line.
        preferred = text.find('：') + 1 if prefer_main_title else 0
        if not preferred or preferred == len(text) or '\n' in text[:preferred] or not (
            width * .5 <= sum(self.metrics.width(c, min(sizes)) for c in text[:preferred]) <= width - 14
        ):
            preferred = 0
        for size in sizes:
            if preferred and sum(self.metrics.width(c, size) for c in text[:preferred]) > width - 14:
                continue
            spans, start, used, failed = [], 0, 0.0, False
            # Keep identifiers such as GRAPH_TABLE intact instead of character wrapping.
            for match in re.finditer(r'(?:[A-Za-z0-9_]+(?:[.-][A-Za-z0-9_]+)*|[^\n])[,.;:!?，。；：！？、）】》”’]*|\n', text):
                token = match.group()
                if token == '\n':
                    spans.append((start, match.start())); start, used = match.end(), 0.0
                    continue
                advance = sum(self.metrics.width(char, size) for char in token)
                if advance > width - 14:
                    failed = True; break
                if used + advance > width - 14 and match.start() > start:
                    spans.append((start, match.start())); start, used = match.start(), 0.0
                used += advance
                if preferred and match.end() == preferred:
                    spans.append((start, preferred)); start, used = preferred, 0.0
            if failed:
                continue
            spans.append((start, len(text)))
            measured = len(spans) * size * 1.4 + 4
            if len(spans) <= max_lines and measured <= height:
                return {'size': size, 'h': measured, 'runs': runs,
                        'line_runs': [slice_runs(runs, a, b) for a, b in spans]}
        if preferred:
            return self.heading(runs, width, height, sizes, max_lines, label, prefer_main_title=False)
        raise ValueError(f'{label}过长：已尝试自动换行及缩小至 {min(sizes)} pt，仍超出 {max_lines} 行的安全区域。请缩短标题或拆分章节；不会截断原文。')

    def new_page(self, title: str | None = None, block: str | None = None, runs=None):
        if len(self.pages) >= 160:
            raise ValueError('幻灯片超过 160 页，请拆分文档后导出。')
        if title is not None:
            self.title = title
            self.section_pages = 0
        display = self.title if self.section_pages == 0 else f"{self.title}（续）"
        self.section_pages += 1
        title_runs = runs if runs is not None else [{"text": display}]
        heading = self.heading(title_runs, 11.67 * 72, 94, (32, 30, 28, 26, 24), 2, '章节标题')
        self.page = {"layout": "content", "title": display, "elements": []}
        self.pages.append(self.page)
        self.add({"kind": "text", "block": block, "x": .84 * 72, "y": .2 * 72, "w": 11.67 * 72,
                  "bold": True, **heading})
        self.y = self.box[1]

    def add(self, element: dict):
        if self.sequence >= 6000:
            raise ValueError('页面对象超过 6,000 个，请拆分文档后导出。')
        self.sequence += 1
        element["name"] = f"tk:{element.get('block') or 'context'}:{self.sequence}"
        self.page["elements"].append(element)

    def room(self, height: float):
        if height > self.box[3] + .01:
            raise ValueError("An indivisible element exceeds the slide body; manual layout required")
        if self.page is None or self.y + height > self.box[1] + self.box[3]:
            self.new_page()

    def paragraph(self, block: dict, runs: list[dict], size: float, code=False):
        if any("asset" in run for run in runs):
            return self.inline_graphics(block, runs, size)
        text = text_of(runs)
        if not text:
            return
        line_height = size * 1.40
        # Account for frame margins and avoid consuming the footer's safety zone.
        lines = self.metrics.lines(text, self.box[2] - 14, size)
        position = 0
        while position < len(lines):
            self.room(line_height + 8)
            available = int((self.box[1] + self.box[3] - self.y - 8) // line_height)
            count = max(1, min(available, len(lines) - position))
            end = lines[position + count - 1][1]
            chunk = slice_runs(runs, lines[position][0], end)
            height = count * line_height + 8
            self.add({"kind": "text", "block": block["id"], "x": self.box[0], "y": self.y, "w": self.box[2], "h": height,
                      "size": size, "code": code, "quote": block.get("quote", False), "runs": chunk})
            self.y += height + 8
            position += count

    def inline_graphics(self, block, runs, size):
        # Lay out native text fragments and embedded formula pictures in reading order.
        x, line_height = self.box[0], size * 1.6
        self.room(line_height)
        for run in runs:
            if "asset" in run:
                asset = self.assets[run["asset"]]
                height = size * 1.25
                width = height * asset["width"] / asset["height"]
                if width > self.box[2]:
                    raise ValueError("Inline formula too wide; use a display formula")
                if x + width > self.box[0] + self.box[2]:
                    self.y += line_height; self.room(line_height); x = self.box[0]
                self.add({"kind": "image", "block": block["id"], "asset": asset["id"], "x": x, "y": self.y + 3, "w": width, "h": height})
                x += width + 4
                continue
            value, start = run.get("text", ""), 0
            while start < len(value):
                available = self.box[0] + self.box[2] - x - 6
                if available < size * 1.2:
                    self.y += line_height; self.room(line_height); x = self.box[0]; continue
                end = self.metrics.lines(value[start:], available, size)[0][1] + start
                fragment = value[start:end]
                width = self.metrics.width(fragment.rstrip('\n'), size) + 4
                self.add({"kind": "text", "block": block["id"], "x": x, "y": self.y, "w": width, "h": line_height,
                          "size": size, "runs": [{**run, "text": fragment}], "nowrap": True})
                x += width; start = end
                if start < len(value) or fragment.endswith('\n'):
                    self.y += line_height; self.room(line_height); x = self.box[0]
        self.y += line_height + 8

    def graphic(self, block, following=None):
        asset = self.assets[block["asset"]]
        ratio = asset["width"] / asset["height"]
        reserve = 0
        if following and following['type'] == 'paragraph' and not any('asset' in run for run in following['runs']):
            size = self.profile['pptx']['body_pt']
            count = len(self.metrics.lines(following.get('prefix', '') + text_of(following['runs']), self.box[2] - 14, size))
            if count <= 3: reserve = count * size * 1.4 + 16
        # Formula images are kept at readable natural size rather than blown up.
        maximum_h = min(self.box[3] - 26 - reserve, 80 if asset["kind"] == "formula" else self.box[3] - 26)
        width, height = min(self.box[2], maximum_h * ratio), min(maximum_h, self.box[2] / ratio)
        # Let a moderate image resize keep its introduction on the same slide.
        # Never crop, distort, or squeeze a graphic into an unreadably small gap.
        if self.page is not None:
            remaining = self.box[1] + self.box[3] - self.y - 26 - reserve
            if min(height * .75, 120) <= remaining < height:
                height = remaining; width = height * ratio
        self.room(height + 26 + reserve)
        self.add({"kind": "image", "block": block["id"], "asset": asset["id"], "x": self.box[0] + (self.box[2] - width) / 2,
                  "y": self.y, "w": width, "h": height})
        self.y += height + 26

    def table(self, block):
        rows = block["rows"]
        if not rows or not rows[0]:
            raise ValueError("Empty table")
        if any(len(row) != len(rows[0]) for row in rows):
            raise ValueError("Inconsistent table columns")
        if any('asset' in run for row in rows for cell in row for run in cell):
            raise ValueError("Table cells with graphics require manual layout in this proof")
        size, count = self.profile["pptx"]["table_pt"], len(rows[0])
        if count > 6:
            raise ValueError("Proof supports at most 6 table columns; no silent shrinking")
        widths = [self.box[2] / count] * count
        if count == 3:
            widths = [self.box[2] * .15, self.box[2] * .40, self.box[2] * .45]
        heights = [max(len(self.metrics.lines(text_of(cell), widths[i] - 20, size)) for i, cell in enumerate(row)) * size * 1.4 + 14 for row in rows]
        if any(h + heights[0] > self.box[3] for h in heights[1:]):
            raise ValueError("A table row exceeds a slide; manual restructuring required")
        start = 1
        while start < len(rows):
            self.room(heights[0] + heights[start])
            end, height = start, heights[0]
            while end < len(rows) and self.y + height + heights[end] <= self.box[1] + self.box[3]:
                height += heights[end]; end += 1
            self.add({"kind": "table", "block": block["id"], "x": self.box[0], "y": self.y, "w": self.box[2], "h": height,
                      "size": size, "rows": [rows[0], *rows[start:end]], "row_indices": [0, *range(start, end)],
                      "widths": widths, "heights": [heights[0], *heights[start:end]]})
            self.y += height + 12
            start = end
        if len(rows) == 1:
            self.room(heights[0])
            self.add({"kind": "table", "block": block["id"], "x": self.box[0], "y": self.y, "w": self.box[2], "h": heights[0],
                      "size": size, "rows": rows, "row_indices": [0], "widths": widths, "heights": heights})
            self.y += heights[0] + 12

    def build(self):
        for index, block in enumerate(self.source["blocks"]):
            kind = block["type"]
            if index == 0 and kind == "heading" and block["level"] == 1:
                title = text_of(block["runs"])
                if any('asset' in run for run in block['runs']):
                    raise ValueError('封面标题暂不支持图片或公式，请移至正文。')
                heading = self.heading(block['runs'], 7 * 72, 160, (36, 34, 32, 30, 28), 3, '封面标题')
                subtitle_y = max(3.85 * 72, 2.35 * 72 + heading['h'] + 12)
                version_y = max(4.97 * 72, subtitle_y + .6 * 72 + 20)
                self.page = {"layout": "cover", "title": title, "elements": []}
                self.pages.append(self.page)
                self.add({"kind": "text", "block": block["id"], "x": .87 * 72, "y": 2.35 * 72, "w": 7 * 72,
                          "bold": True, **heading})
                self.add({"kind": "text", "block": None, "x": .87 * 72, "y": subtitle_y, "w": 7 * 72, "h": .6 * 72,
                          "size": 16, "runs": [{"text": self.profile['sample_metadata']['subtitle']}]})
                self.add({"kind": "text", "block": None, "x": .87 * 72, "y": version_y, "w": 5.55 * 72, "h": .6 * 72,
                          "size": 14, "runs": [{"text": self.profile['sample_metadata']['version']}]})
                self.page = None; self.title = "概览"
                self.section_pages = 0
            elif kind == "heading" and block["level"] <= 2:
                if any('asset' in run for run in block['runs']):
                    raise ValueError('章节标题暂不支持图片或公式，请移至正文。')
                self.new_page(text_of(block["runs"]), block["id"], block['runs'])
            elif kind in ("paragraph", "heading"):
                if block.get('list') and (index == 0 or self.source['blocks'][index - 1].get('list') != block['list']):
                    # Keep a short contiguous list together instead of orphaning its last item.
                    height = 0
                    for candidate in self.source['blocks'][index:]:
                        if candidate.get('list') != block['list'] or any('asset' in run for run in candidate.get('runs', [])): break
                        size = self.profile['pptx']['body_pt']
                        height += len(self.metrics.lines(candidate.get('prefix', '') + text_of(candidate['runs']), self.box[2] - 14, size)) * size * 1.4 + 16
                    if 0 < height <= self.box[3]: self.room(height)
                runs = [{"text": block["prefix"]}] + block['runs'] if block.get('prefix') else block['runs']
                if kind == "heading":
                    runs = [{**run, "bold": True} for run in runs]
                self.paragraph(block, runs, self.profile["pptx"]["body_pt"])
            elif kind == "code":
                if block.get('language') == 'mermaid':
                    if self.page is None: self.new_page()
                    self.page.setdefault('notes', []).append({'block': block['id'], 'text': block['text']})
                else:
                    self.paragraph(block, [{"text": block["text"]}], self.profile["pptx"]["code_pt"], code=True)
            elif kind == "graphic":
                next_index = index + 1
                if next_index < len(self.source['blocks']) and self.source['blocks'][next_index].get('language') == 'mermaid': next_index += 1
                following = self.source['blocks'][next_index] if next_index < len(self.source['blocks']) else None
                self.graphic(block, following)
            elif kind == "table":
                self.table(block)
            elif kind == "break":
                self.page = None
            elif kind == "rule":
                self.room(12)
                self.add({"kind": "rule", "block": block["id"], "x": self.box[0], "y": self.y, "w": self.box[2], "h": 2})
                self.y += 12
            else:
                raise ValueError(f"Unsupported block: {kind}")
        return self.pages


def ppt_font(run, size, profile, bold=False, code=False):
    run.font.size = Pt(size)
    run.font.name = profile['fonts']['code' if code else 'latin']
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(profile['colors']['text'])
    rpr = run._r.get_or_add_rPr()
    ea = PptElement('a:ea'); ea.set('typeface', profile['fonts']['east_asia']); rpr.append(ea)


def ppt_text(frame, runs, size, profile, *, bold=False, code=False, nowrap=False):
    frame.clear(); frame.word_wrap = not nowrap
    frame.auto_size = MSO_AUTO_SIZE.NONE
    frame.margin_left = frame.margin_right = Pt(3)
    frame.margin_top = frame.margin_bottom = Pt(1)
    frame.vertical_anchor = MSO_ANCHOR.TOP
    paragraph = frame.paragraphs[0]
    paragraph.line_spacing = Pt(size * 1.4)
    paragraph.space_before = paragraph.space_after = Pt(0)
    for item in runs:
        if 'asset' in item:
            raise ValueError('Unplaced inline asset')
        pieces = item.get('text', '').split('\n')
        for i, text in enumerate(pieces):
            if i:
                paragraph.add_line_break()
            if text:
                run = paragraph.add_run(); run.text = text
                ppt_font(run, size, profile, bold or item.get('bold', False), code or item.get('code', False))
                run.font.italic = item.get('italic', False)
                if item.get('href'):
                    run.hyperlink.address = item['href']


def clean_presentation(prs):
    # Drop all example slides and their relationships, not just the visible IDs.
    for item in list(prs.slides._sldIdLst):
        prs.part.drop_rel(item.rId); prs.slides._sldIdLst.remove(item)
    for owner in [*prs.slide_masters, *prs.slide_layouts]:
        for shape in list(owner.shapes):
            if shape.is_placeholder:
                # Clear demonstration strings, but retain template styles/geometry.
                if shape.has_text_frame:
                    for p in shape.text_frame.paragraphs:
                        for child in list(p._p):
                            if child.tag.rsplit('}', 1)[-1] != 'pPr':
                                p._p.remove(child)
                if shape.placeholder_format.type in (13, 15, 16):  # slide number, footer, date
                    shape._element.getparent().remove(shape._element)


def write_pptx(source, profile, pages, directory):
    prs = Presentation(ROOT / profile['pptx']['path'])
    expected = profile['pptx']['expected_slide_inches']
    if abs(prs.slide_width / 914400 - expected[0]) > .001 or abs(prs.slide_height / 914400 - expected[1]) > .001:
        raise ValueError('Template slide size changed; adapter must be reviewed')
    clean_presentation(prs)
    layouts = {layout.name: layout for layout in prs.slide_layouts}
    assets = {asset['id']: asset for asset in source['assets']}
    for number, page in enumerate(pages, 1):
        layout = layouts[profile['pptx']['layouts'][page['layout']]]
        slide = prs.slides.add_slide(layout)
        for shape in list(slide.placeholders):
            shape._element.getparent().remove(shape._element)
        for element in page['elements']:
            x, y, w, h = [Pt(element[key]) for key in ('x', 'y', 'w', 'h')]
            if element['kind'] == 'text':
                shape = slide.shapes.add_textbox(x, y, w, h)
                runs = element['runs']
                if 'line_runs' in element:
                    runs = []
                    for i, line in enumerate(element['line_runs']):
                        if i: runs.append({'text': '\n'})
                        runs.extend(line)
                ppt_text(shape.text_frame, runs, element['size'], profile, bold=element.get('bold', False), code=element.get('code', False), nowrap='line_runs' in element or element.get('nowrap', False))
                if element.get('code') or element.get('quote'):
                    shape.fill.solid(); shape.fill.fore_color.rgb = RGBColor.from_string(profile['colors']['table_alternate'])
            elif element['kind'] == 'image':
                asset = assets[element['asset']]
                shape = slide.shapes.add_picture(str(safe_asset(directory, asset['file'])), x, y, width=w, height=h)
                shape._element.nvPicPr.cNvPr.set('descr', asset['alt'])
            elif element['kind'] == 'table':
                shape = slide.shapes.add_table(len(element['rows']), len(element['widths']), x, y, w, h)
                table = shape.table
                for i, width in enumerate(element['widths']): table.columns[i].width = Pt(width)
                for i, row in enumerate(element['rows']):
                    table.rows[i].height = Pt(element['heights'][i])
                    for j, runs in enumerate(row):
                        cell = table.cell(i, j)
                        ppt_text(cell.text_frame, runs, element['size'], profile, bold=i == 0)
                        cell.margin_top = Pt(6); cell.margin_bottom = Pt(4)
                        cell.margin_left = cell.margin_right = Pt(8)
                        cell.fill.solid(); cell.fill.fore_color.rgb = RGBColor.from_string(profile['colors']['table_header'] if i == 0 else profile['colors']['table_alternate'] if i % 2 else 'FFFFFF')
                        if i == 0:
                            for run in cell.text_frame.paragraphs[0].runs: run.font.color.rgb = RGBColor(255, 255, 255)
            elif element['kind'] == 'rule':
                shape = slide.shapes.add_shape(1, x, y, w, h)
                shape.fill.solid(); shape.fill.fore_color.rgb = RGBColor.from_string(profile['colors']['accent'])
                shape.line.fill.background()
            else:
                raise ValueError(f"Unsupported layout element: {element['kind']}")
            shape.name = element['name']
        for x, width, text in [(.84, .35, str(number)), (1.23, 7.0, profile['pptx']['copyright']), (9, 3.6, profile['sample_metadata']['footer'])]:
            box = slide.shapes.add_textbox(Inches(x), Inches(7.04), Inches(width), Inches(.28))
            box.name = 'tk:footer'
            ppt_text(box.text_frame, [{'text': text}], 8, profile)
        if page.get('notes'):
            slide.notes_slide.notes_text_frame.text = '\n\n'.join(f'Markdown source [{item["block"]}]\n{item["text"]}' for item in page['notes'])
    prs.core_properties.title = pages[0]['title']
    prs.core_properties.author = ''; prs.core_properties.last_modified_by = ''
    prs.core_properties.subject = profile.get('document_subject', 'Template adaptation proof — not delivery-approved')
    prs.core_properties.keywords = ''; prs.core_properties.comments = ''
    prs.save(directory / 'aibs-markdown-proof.pptx')


def word_font(run, profile, size=None, code=False):
    # DOCX keeps each original template style's Latin font (Title/Heading/Normal).
    if code: run.font.name = profile['fonts']['code']
    if size: run.font.size = WordPt(size)
    rfonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    rfonts.set(qn('w:eastAsia'), profile['fonts']['east_asia'])


def word_runs(paragraph, runs, profile, assets, directory, code=False):
    for item in runs:
        run = paragraph.add_run()
        if 'asset' in item:
            asset = assets[item['asset']]
            graphic = run.add_picture(str(safe_asset(directory, asset['file'])), height=WordPt(15))
            graphic._inline.docPr.set('descr', asset['alt'])
            continue
        run.text = item.get('text', '')
        if 'bold' in item: run.bold = item['bold']
        if 'italic' in item: run.italic = item['italic']
        word_font(run, profile, 9 if code else None, code or item.get('code', False))
        if item.get('href'):
            link = OxmlElement('w:hyperlink')
            link.set(qn('r:id'), paragraph.part.relate_to(item['href'], REL.HYPERLINK, is_external=True))
            run.font.color.rgb = WordColor.from_string('2C5967'); run.font.underline = True
            link.append(run._r); paragraph._p.append(link)


def word_field(paragraph, instruction, cached):
    field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), instruction)
    run = OxmlElement('w:r'); text = OxmlElement('w:t'); text.text = cached
    run.append(text); field.append(run); paragraph._p.append(field)


def write_docx(source, profile, directory):
    doc = Document(ROOT / profile['docx']['path'])
    if len(doc.sections) != 2:
        raise ValueError('Word section structure changed; adapter must be reviewed')
    cover_section = deepcopy(doc.sections[0]._sectPr)
    body_section = deepcopy(doc.sections[1]._sectPr)
    body = doc.element.body
    for element in list(body): body.remove(element)
    body.append(body_section)
    assets = {asset['id']: asset for asset in source['assets']}
    for name in profile['docx']['styles'].values():
        if name not in doc.styles: raise ValueError(f'Missing Word style: {name}')
    # Keep template paragraph design; apply an explicit Chinese font proposal.
    for name in profile['docx']['styles'].values():
        style = doc.styles[name]
        rpr = style.element.get_or_add_rPr()
        fonts = rpr.find(qn('w:rFonts'))
        if fonts is None: fonts = OxmlElement('w:rFonts'); rpr.insert(0, fonts)
        fonts.set(qn('w:eastAsia'), profile['fonts']['east_asia'])
    blocks = source['blocks']
    has_title = bool(blocks and blocks[0]['type'] == 'heading' and blocks[0]['level'] == 1)
    title = text_of(blocks[0]['runs']) if has_title else 'Markdown 文档'
    paragraph = doc.add_paragraph(style='Title')
    word_runs(paragraph, blocks[0]['runs'] if has_title else [{'text': title}], profile, assets, directory)
    for text, style in [(profile['sample_metadata']['subtitle'], 'Subtitle'), (profile['sample_metadata']['version'], 'Cover date and Info'), (profile['docx']['copyright'], 'Cover date and Info')]:
        paragraph = doc.add_paragraph(style=style); word_runs(paragraph, [{'text': text}], profile, assets, directory)
    boundary = doc.add_paragraph()
    boundary._p.get_or_add_pPr().append(cover_section)
    # TOC is deliberately not fabricated: pagination must be verified in Office.
    for block in blocks[1:] if has_title else blocks:
        kind = block['type']
        if kind in ('paragraph', 'heading', 'code'):
            style = 'Heading ' + str(min(3, block['level'])) if kind == 'heading' else 'Code' if kind == 'code' else 'Normal'
            paragraph = doc.add_paragraph(style=style)
            if kind == 'code':
                paragraph.paragraph_format.keep_with_next = False
                paragraph.paragraph_format.keep_together = False
                paragraph.paragraph_format.space_before = WordPt(6)
                paragraph.paragraph_format.space_after = WordPt(8)
                shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), profile['colors']['table_alternate']); paragraph._p.get_or_add_pPr().append(shade)
            if block.get('quote'):
                paragraph.paragraph_format.left_indent = WordInches(.2)
            if block.get('prefix'):
                word_runs(paragraph, [{'text': block['prefix']}], profile, assets, directory)
            runs = [{'text': block['text']}] if kind == 'code' else block['runs']
            word_runs(paragraph, runs, profile, assets, directory, code=kind == 'code')
        elif kind == 'graphic':
            asset = assets[block['asset']]
            section = doc.sections[-1]
            width = (section.page_width - section.left_margin - section.right_margin) / 914400
            max_height = 5.8 if asset['kind'] != 'formula' else .7
            width = min(width, max_height * asset['width'] / asset['height'])
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.keep_together = True
            picture = paragraph.add_run().add_picture(str(safe_asset(directory, asset['file'])), width=WordInches(width))
            picture._inline.docPr.set('descr', asset['alt'])
        elif kind == 'table':
            table = doc.add_table(rows=0, cols=len(block['rows'][0]))
            table.autofit = False
            table.style = 'Table Grid'
            section = doc.sections[-1]
            width = (section.page_width - section.left_margin - section.right_margin) / 914400
            for column in table.columns: column.width = WordInches(width / len(table.columns))
            for index, row in enumerate(block['rows']):
                cells = table.add_row().cells
                props = table.rows[-1]._tr.get_or_add_trPr()
                props.append(OxmlElement('w:cantSplit'))
                if index == 0: props.append(OxmlElement('w:tblHeader'))
                for cell, runs in zip(cells, row):
                    word_runs(cell.paragraphs[0], runs, profile, assets, directory)
                    shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), profile['colors']['table_header'] if index == 0 else profile['colors']['table_alternate'] if index % 2 else 'FFFFFF')
                    cell._tc.get_or_add_tcPr().append(shade)
                    if index == 0:
                        for run in cell.paragraphs[0].runs:
                            run.bold = True; run.font.color.rgb = WordColor(255, 255, 255)
            doc.add_paragraph()
        elif kind == 'rule':
            paragraph = doc.add_paragraph()
            border = OxmlElement('w:pBdr'); bottom = OxmlElement('w:bottom')
            for key, value in {'val': 'single', 'sz': '8', 'color': profile['colors']['accent']}.items(): bottom.set(qn('w:' + key), value)
            border.append(bottom); paragraph._p.get_or_add_pPr().append(border)
        elif kind == 'break':
            # Slide directives do not force unrelated Word page breaks.
            continue
        else:
            raise ValueError(f'Unsupported Word block: {kind}')
    for section in doc.sections:
        for footer in (section.footer, section.first_page_footer, section.even_page_footer):
            footer.is_linked_to_previous = False
            for child in list(footer._element): footer._element.remove(child)
            paragraph = footer.add_paragraph()
            word_runs(paragraph, [{'text': profile['docx']['copyright'] + ' / ' + profile['sample_metadata']['footer'] + ' / '}], profile, assets, directory)
            for run in paragraph.runs: run.font.size = WordPt(8)
            word_field(paragraph, 'PAGE', '1')
    update = doc.settings.element.find(qn('w:updateFields'))
    if update is None: update = OxmlElement('w:updateFields'); doc.settings.element.append(update)
    update.set(qn('w:val'), 'true')
    doc.core_properties.title = title
    doc.core_properties.author = ''; doc.core_properties.last_modified_by = ''
    doc.core_properties.subject = profile.get('document_subject', 'Template adaptation proof — not delivery-approved')
    doc.core_properties.keywords = ''; doc.core_properties.comments = ''
    # Remove stale example comments/embeddings/custom data relationships after body replacement.
    for rel in list(doc.part.rels.values()):
        if rel.reltype.rsplit('/', 1)[-1] in ('comments', 'customXml', 'chart', 'oleObject', 'package'):
            doc.part.drop_rel(rel.rId)
    doc.save(directory / 'aibs-markdown-proof.docx')


def write_preview(source, profile, pages, directory):
    """An explicitly approximate HTML layout proof, never labelled an Office render."""
    template = Presentation(ROOT / profile['pptx']['path'])
    layouts = {layout.name: layout for layout in template.slide_layouts}
    assets = {asset['id']: asset for asset in source['assets']}
    def styled_runs(runs):
        out = []
        for run in runs:
            text = html.escape(run.get('text', ''))
            if run.get('bold'): text = f'<b>{text}</b>'
            if run.get('italic'): text = f'<i>{text}</i>'
            if run.get('code'): text = f'<code>{text}</code>'
            out.append(text)
        return ''.join(out)
    markup = ['<!doctype html><meta charset="utf-8"><title>AIBS 排版检查样稿</title>',
              '<style>body{background:#e7e5e4;margin:24px;color:#312d2a;font-family:"Noto Sans CJK SC",Arial,sans-serif} .page{position:relative;width:960pt;height:540pt;margin:24px 0;background:#fcfbfa;overflow:hidden} .item{position:absolute;box-sizing:border-box} .text{padding:1pt 3pt;white-space:pre-wrap;overflow:hidden} table{border-collapse:collapse;table-layout:fixed} td{padding:6pt 8pt;vertical-align:top;overflow:hidden} .warning{max-width:1000px;background:#fff4dc;padding:16px} code{font-family:Consolas,monospace}</style>',
              '<div class="warning"><b>仅供排版检查：HTML 近似预览，不是 Office 文件渲染。</b><br>原模板装饰图＋生成布局；字体替换、Word 分页和 Office 兼容性仍须人工验收。文字与表格在 PPTX 中不是整页截图。<br>浏览器不能显示的 WMF 母版装饰用虚线框标记；PPTX 中保留原始图形。</div>']
    for number, page in enumerate(pages, 1):
        markup.append(f'<section class="page" data-page="{number}">')
        layout = layouts[profile['pptx']['layouts'][page['layout']]]
        owners = [layout.slide_master, layout] if layout._element.get('showMasterSp') != '0' else [layout]
        for owner_index, owner in enumerate(owners):
            for i, shape in enumerate(owner.shapes):
                if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                    image = shape.image
                    if image.ext not in ('png', 'jpg', 'jpeg', 'gif', 'webp', 'svg'):
                        markup.append(f'<div class="item" data-preview-placeholder="{image.ext}" title="{image.ext} decoration retained in PPTX, not rendered here" style="left:{shape.left/12700}pt;top:{shape.top/12700}pt;width:{shape.width/12700}pt;height:{shape.height/12700}pt;border:1px dashed #c74634;font-size:6pt">{image.ext}</div>')
                        continue
                    name = f'assets/template-{page["layout"]}-{owner_index}-{i}.{image.ext}'
                    (directory / name).write_bytes(image.blob)
                    markup.append(f'<img class="item" src="{name}" alt="" style="left:{shape.left/12700}pt;top:{shape.top/12700}pt;width:{shape.width/12700}pt;height:{shape.height/12700}pt">')
        for element in page['elements']:
            geometry = ';'.join(f'{css}:{element[key]}pt' for css, key in [('left', 'x'), ('top', 'y'), ('width', 'w'), ('height', 'h')])
            kind = element['kind']
            if kind == 'text':
                extra = 'font-family:Consolas,"Noto Sans CJK SC",monospace;' if element.get('code') else ''
                extra += 'font-weight:bold;' if element.get('bold') else ''
                extra += 'background:#f4f1ed;' if element.get('code') or element.get('quote') else ''
                content = ''.join(f'<div style="white-space:pre;min-height:1.4em">{styled_runs(line)}</div>' for line in element['line_runs']) if 'line_runs' in element else styled_runs(element['runs'])
                markup.append(f'<div class="item text" data-object="{element["name"]}" style="{geometry};font-size:{element["size"]}pt;line-height:1.4;{extra}">{content}</div>')
            elif kind == 'image':
                asset = assets[element['asset']]
                markup.append(f'<img class="item" src="{asset["file"]}" alt="{html.escape(asset["alt"], quote=True)}" style="{geometry}">')
            elif kind == 'table':
                markup.append(f'<table class="item" data-object="{element["name"]}" style="{geometry};font-size:{element["size"]}pt;line-height:1.4"><colgroup>')
                markup.extend(f'<col style="width:{width}pt">' for width in element['widths'])
                markup.append('</colgroup>')
                for i, row in enumerate(element['rows']):
                    color = 'background:#312d2a;color:white;font-weight:bold' if i == 0 else 'background:#f4f1ed' if i % 2 else 'background:white'
                    markup.append(f'<tr style="height:{element["heights"][i]}pt;{color}">')
                    markup.extend(f'<td>{styled_runs(cell)}</td>' for cell in row)
                    markup.append('</tr>')
                markup.append('</table>')
            else:
                markup.append(f'<div class="item" style="{geometry};background:#c74634"></div>')
        markup.append(f'<div class="item" style="left:60pt;top:507pt;font-size:8pt">{number}　{html.escape(profile["pptx"]["copyright"])}　 ·　{profile["sample_metadata"]["footer"]}</div></section>')
    (directory / 'ppt-layout-preview.html').write_text('\n'.join(markup), encoding='utf-8')


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('directory', type=Path)
    parser.add_argument('--measure-font', type=Path, required=True, help='Existing local font file for conservative layout metrics; not embedded')
    args = parser.parse_args()
    directory = args.directory.resolve()
    source = json.loads((directory / 'source.json').read_text())
    if source.get('schema') != 1: raise ValueError('Unsupported IR schema')
    if hashlib.sha256((directory / 'source.md').read_bytes()).hexdigest() != source['source_sha256']:
        raise ValueError('Source snapshot hash mismatch')
    for name in ('aibs-markdown-proof.pptx', 'aibs-markdown-proof.docx'):
        if (directory / name).exists(): raise ValueError(f'Refusing to overwrite {name}')
    profile = load_template_profile()
    for asset in source['assets']: safe_asset(directory, asset['file'])
    pages = Planner(source, profile, Metrics(args.measure_font)).build()
    write_pptx(source, profile, pages, directory)
    write_docx(source, profile, directory)
    write_preview(source, profile, pages, directory)
    report = {
        'schema': 1, 'status': 'generated-unvalidated', 'delivery_approved': False,
        'source_sha256': source['source_sha256'], 'profile': profile, 'pages': pages,
        'measurement_font': str(args.measure_font),
        'template_sha256': {kind: hashlib.sha256((ROOT / profile[kind]['path']).read_bytes()).hexdigest() for kind in ('pptx', 'docx')},
        'warnings': source['warnings'] + ['Layout measurement is approximate, not Microsoft Office rendering.', 'WMF template decorations are preserved in PPTX but shown as labelled placeholders in HTML.', 'Word TOC is intentionally deferred; no stale cached TOC or STYLEREF is carried over.', 'Word lists have editable literal markers in this proof; native list numbering is deferred.'],
        'native_office_acceptance': {'macOS': 'not_run', 'Windows': 'not_run'},
    }
    (directory / 'layout-plan.json').write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps({'slides': len(pages), 'directory': str(directory), 'delivery_approved': False}))


if __name__ == '__main__': main()
