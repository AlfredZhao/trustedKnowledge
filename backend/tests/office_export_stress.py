"""Opt-in bounded-worker load check; synthetic screenshots, no services or network.

Run with the backend dependencies/font installed: python -m tests.office_export_stress
"""
import base64
import io
import json
from pathlib import Path
import random
import subprocess
import sys
import tempfile
import time
from zipfile import ZipFile

from PIL import Image
from pptx import Presentation
from docx import Document

from app.schemas.office import OfficeRequest
from app.schemas.office_limits import OfficeLimits


def main():
    font = next(p for p in [Path('/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc'), Path('/usr/share/fonts/google-noto-cjk/NotoSansCJK-Regular.ttc')] if p.is_file())
    policy = OfficeLimits()
    results = []
    # Unique RGB screenshot-like images, not duplicated bytes that Office deduplicates.
    for count, strip_height in ((33, 180), (64, 90)):
        source = {'assets': [], 'blocks': [{'id': 'title', 'type': 'heading', 'level': 1, 'runs': [{'text': '揭开 Oracle 属性图的神秘面纱：给 DBA 的 GRAPH_TABLE 内部机制指南'}]}]}
        for i in range(count):
            image = Image.new('RGB', (1600, 900), 'white')
            image.paste(Image.frombytes('RGB', (1600, strip_height), random.Random(i).randbytes(1600 * strip_height * 3)))
            buffer = io.BytesIO(); image.save(buffer, 'PNG')
            source['assets'].append({'id': f'a{i}', 'kind': 'image', 'alt': f'synthetic screenshot {i}', 'width': 1600, 'height': 900, 'data': base64.b64encode(buffer.getvalue()).decode()})
            source['blocks'].append({'id': f'b{i}', 'type': 'graphic', 'asset': f'a{i}'})
        encoded = sum(len(a['data']) for a in source['assets'])
        assert 24_000_000 < encoded <= policy.max_total_bytes
        assert 40_000_000 < count * 1600 * 900 <= policy.max_total_pixels
        payload = OfficeRequest.model_validate({'source': source})
        request = payload.model_dump_json()
        del source, payload
        for kind in ('pptx', 'docx'):
            with tempfile.TemporaryDirectory() as name:
                directory = Path(name)
                (directory / 'request.json').write_text(request)
                (directory / 'limits.json').write_text(policy.model_dump_json())
                command = [sys.executable, '-m', 'app.services.office.worker', name, kind, str(font)]
                stats = directory / 'resource.txt'
                if Path('/usr/bin/time').exists():
                    command = ['/usr/bin/time', '-f', '%M', '-o', str(stats), *command]
                start = time.monotonic()
                result = subprocess.run(command, capture_output=True, timeout=90)
                elapsed = time.monotonic() - start
                assert result.returncode == 0, (directory / 'error.json').read_text() if (directory / 'error.json').exists() else result.stderr.decode()
                file = directory / f'aibs-markdown-proof.{kind}'
                assert file.stat().st_size <= policy.max_output_bytes
                with ZipFile(file) as archive: assert archive.testzip() is None
                if kind == 'pptx':
                    document = Presentation(file)
                    actual = sum(1 for slide in document.slides for shape in slide.shapes if shape.name.startswith('tk:b') and shape.shape_type == 13)
                else:
                    actual = len(Document(file).inline_shapes)
                assert actual == count, (actual, count)
                row = {'images': count, 'encoded_bytes': encoded, 'pixels': count * 1600 * 900, 'format': kind, 'output_bytes': file.stat().st_size, 'seconds': round(elapsed, 2), 'worker_peak_rss_kib': int(stats.read_text().strip()) if stats.exists() else None}
                results.append(row); print(json.dumps(row), flush=True)
    print(json.dumps({'status': 'passed', 'scope': 'synthetic single-job load; not concurrency or native Office acceptance', 'results': results}))


if __name__ == '__main__': main()
